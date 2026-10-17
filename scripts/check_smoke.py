#!/usr/bin/env python3
"""Basic I/O checks only, not visual QA or official OOXML validation."""
from __future__ import annotations
import json
import sys
import zipfile
from pathlib import Path

def main() -> None:
    if len(sys.argv) != 3 or sys.argv[1] not in {'pptx', 'pdf'}:
        raise SystemExit('Usage: check_smoke.py pptx|pdf FILE')
    mode, path = sys.argv[1], Path(sys.argv[2])
    report_path = path.parent / 'smoke-report.json'
    report = json.loads(report_path.read_text()) if report_path.exists() else {}
    if mode == 'pptx':
        from pptx import Presentation
        pres = Presentation(path)
        if len(pres.slides) != 2:
            raise RuntimeError('Expected two slides')
        for n, slide in enumerate(pres.slides, 1):
            text = slide.notes_slide.notes_text_frame.text
            if f'SMOKE_NOTE_{n}' not in text:
                raise RuntimeError(f'Notes missing on slide {n}')
        with zipfile.ZipFile(path) as z:
            if z.testzip() is not None:
                raise RuntimeError('ZIP integrity test failed')
            if 'ppt/charts/chart1.xml' not in z.namelist():
                raise RuntimeError('Native chart missing')
        report.update(slides=2, notes_checked=True, native_chart_part=True, zip_integrity=True)
    else:
        from pypdf import PdfReader
        if len(PdfReader(path).pages) != 2:
            raise RuntimeError('Expected two PDF pages')
        if len(list(path.parent.glob('slide-*.png'))) != 2:
            raise RuntimeError('Expected two rendered PNG files')
        text = (path.parent / 'smoke.txt').read_text()
        if '日本語' not in text:
            raise RuntimeError('Japanese text extraction failed; inspect font/rendering setup')
        report.update(pdf_pages=2, png_pages=2, japanese_text_extractable=True,
                      visual_qa='not performed by this script',
                      powerpoint_application_test='not performed',
                      official_skill_validation='must run in Claude production workflow')
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(f'{mode}: basic checks passed')

if __name__ == '__main__':
    main()
